"""Bulk approved-answer import in any format (Phase 8 C).

Extracts Q&A pairs from CSV / JSON / XLSX (and a best-effort two-column DOCX/MD
table) and routes each pair through approve_one — so it trains the library with
dedup/versioning, exactly like every other accepted answer. Reports per-file
counts + reasons; never bypasses the flywheel.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from .flywheel import approve_one

_QKEYS = ("question", "q", "prompt")
_AKEYS = ("answer", "a", "response")


def _pick(d: dict, keys) -> str:
    for k in d:
        if str(k).strip().lower() in keys:
            return str(d[k] or "").strip()
    return ""


def extract_pairs(filename: str, data: bytes) -> list[tuple[str, str]]:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        rows = csv.DictReader(io.StringIO(data.decode("utf-8", errors="replace")))
        return [(q, a) for r in rows if (q := _pick(r, _QKEYS)) and (a := _pick(r, _AKEYS))]
    if ext == ".json":
        obj = json.loads(data.decode("utf-8", errors="replace"))
        items = obj if isinstance(obj, list) else obj.get("qa") or obj.get("pairs") or []
        out = []
        for it in items:
            if isinstance(it, dict):
                q, a = _pick(it, _QKEYS), _pick(it, _AKEYS)
                if q and a:
                    out.append((q, a))
        return out
    if ext in {".xlsx", ".xlsm"}:
        return _from_xlsx(data)
    if ext in {".md", ".markdown", ".txt"}:
        return _from_md_table(data.decode("utf-8", errors="replace"))
    if ext == ".docx":
        return _from_docx(data)
    raise ValueError(f"unsupported Q&A format '{ext}'")


def _from_xlsx(data: bytes) -> list[tuple[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c or "").strip().lower() for c in rows[0]]
    qi = next((i for i, h in enumerate(header) if h in _QKEYS), 0)
    ai = next((i for i, h in enumerate(header) if h in _AKEYS), 1)
    start = 1 if any(h in _QKEYS or h in _AKEYS for h in header) else 0
    out = []
    for row in rows[start:]:
        q = str(row[qi]).strip() if qi < len(row) and row[qi] else ""
        a = str(row[ai]).strip() if ai < len(row) and row[ai] else ""
        if q and a:
            out.append((q, a))
    return out


def _from_md_table(text: str) -> list[tuple[str, str]]:
    out = []
    for line in text.splitlines():
        if "|" not in line:
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 2:
            continue
        q, a = cells[0], cells[1]
        if not q or not a or set(q) <= set("-: ") or q.lower() in _QKEYS:
            continue  # skip separators / header
        out.append((q, a))
    return out


def _from_docx(data: bytes) -> list[tuple[str, str]]:
    import docx as _docx

    document = _docx.Document(io.BytesIO(data))
    out = []
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0] and cells[1] and cells[0].lower() not in _QKEYS:
                out.append((cells[0], cells[1]))
    return out


def import_qa(items, qa_path, approved_by: str = "import", tags=None) -> dict:
    """items: iterable of (filename, bytes). Returns counts + per-file reasons."""
    imported = 0
    per_file = []
    for name, data in items:
        try:
            pairs = extract_pairs(name, data)
        except Exception as exc:  # noqa: BLE001
            per_file.append({"name": name, "ok": False, "reason": str(exc)})
            continue
        for q, a in pairs:
            approve_one(q, a, qa_path, approved_by=approved_by, tags=tags)
            imported += 1
        per_file.append({"name": name, "ok": True, "pairs": len(pairs)})
    return {"imported": imported, "files": per_file}
